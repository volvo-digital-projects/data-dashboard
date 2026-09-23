from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


YEARS = {"2023", "2024", "2025", "2026"}
YEAR_ORDER = tuple(sorted(YEARS))


def normalise(value: Any) -> str:
    return str(value or "").strip().replace(" ", "")


def sent_year(value: Any) -> str | None:
    if isinstance(value, datetime):
        year = str(value.year)
    else:
        year = str(value or "").strip()[:4]
    return year if year in YEARS else None


def response_date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        try:
            return datetime.fromisoformat(text[:10])
        except ValueError:
            return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sent", required=True, type=Path)
    parser.add_argument("--voc", required=True, type=Path)
    parser.add_argument("--analysis", required=True, type=Path)
    parser.add_argument(
        "--sent-output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "app" / "data" / "voc-sent.json",
    )
    args = parser.parse_args()

    payload = json.loads(args.analysis.read_text(encoding="utf-8"))
    employee_index: dict[tuple[str, str], dict[str, Any]] = {}
    for showroom in payload["showrooms"].values():
        showroom_name = normalise(showroom["showroom"]).removeprefix("볼보")
        for employee in showroom["employees"]:
            employee_index[(showroom_name, normalise(employee["name"]))] = employee

    workbook = openpyxl.load_workbook(args.sent, read_only=True, data_only=True)
    worksheet = workbook["Raw Data"] if "Raw Data" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    headers = {
        str(value or "").strip(): index
        for index, value in enumerate(
            next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))
        )
        if value
    }
    date_index = headers["전시장 방문일시"]
    showroom_index = headers["소속 전시장명"]
    employee_index_column = headers["담당 영업직원"]

    counts: Counter[tuple[str, str, str]] = Counter()
    national_counts: Counter[str] = Counter()
    showroom_counts: Counter[tuple[str, str]] = Counter()
    latest_sent_at: datetime | None = None
    for row in worksheet.iter_rows(min_row=3, values_only=True):
        year = sent_year(row[date_index])
        showroom = normalise(row[showroom_index])
        employee = normalise(row[employee_index_column])
        if not year:
            continue
        national_counts[year] += 1
        if showroom:
            showroom_counts[(showroom, year)] += 1
        if showroom and employee:
            counts[(showroom, employee, year)] += 1
        parsed_sent_at = response_date(row[date_index])
        if parsed_sent_at is not None and (
            latest_sent_at is None or parsed_sent_at > latest_sent_at
        ):
            latest_sent_at = parsed_sent_at

    if latest_sent_at is None:
        raise RuntimeError("VOC 발송 원본에서 유효한 전시장 방문일시를 찾지 못했습니다.")

    showroom_cdsid = {
        normalise(showroom["showroom"]).removeprefix("볼보"): cdsid
        for cdsid, showroom in payload["showrooms"].items()
    }
    source_showrooms = {showroom for showroom, _ in showroom_counts}
    if source_showrooms != set(showroom_cdsid):
        raise RuntimeError(
            "VOC 발송 전시장 매핑 불일치: "
            f"누락={sorted(set(showroom_cdsid) - source_showrooms)}, "
            f"초과={sorted(source_showrooms - set(showroom_cdsid))}"
        )

    matched_rows = 0
    matched_employees: set[tuple[str, str]] = set()
    for (showroom, employee, year), sent in counts.items():
        target = employee_index.get((showroom, employee))
        if target is None:
            continue
        metric = target["years"].setdefault(year, {"responses": 0, "scoreSum": 0})
        metric["sent"] = sent
        matched_rows += sent
        matched_employees.add((showroom, employee))

    for employee in employee_index.values():
        for metric in employee["years"].values():
            metric.setdefault("sent", 0)

    voc_workbook = openpyxl.load_workbook(args.voc, read_only=True, data_only=True)
    latest_response: dict[tuple[str, str], datetime] = {}
    for year in sorted(YEARS):
        voc_sheet = voc_workbook[year]
        voc_headers = {
            str(value or "").strip().replace("\n", ""): index
            for index, value in enumerate(
                next(voc_sheet.iter_rows(min_row=2, max_row=2, values_only=True))
            )
            if value
        }
        visit_index = voc_headers["전시장 방문일"]
        voc_showroom_index = voc_headers.get("방문 전시장", voc_headers.get("방문"))
        voc_employee_index = voc_headers["상담영업직원"]
        for row in voc_sheet.iter_rows(min_row=3, values_only=True):
            visited_at = response_date(row[visit_index])
            showroom = normalise(row[voc_showroom_index]) if voc_showroom_index is not None else ""
            employee = normalise(row[voc_employee_index])
            key = (showroom, employee)
            if visited_at and key in employee_index and (
                key not in latest_response or visited_at > latest_response[key]
            ):
                latest_response[key] = visited_at

    for key, employee in employee_index.items():
        latest = latest_response.get(key)
        if latest:
            employee["latestResponseDate"] = latest.date().isoformat()
        else:
            employee.pop("latestResponseDate", None)

    sent_through = latest_sent_at.date().isoformat()
    payload["source"]["sentWorkbook"] = args.sent.name
    payload["source"]["sentThrough"] = sent_through
    payload["source"]["latestResponseThrough"] = payload["source"]["vocThrough"]
    args.analysis.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    sent_payload = {
        "updatedThrough": sent_through,
        "sourceWorkbook": args.sent.name,
        "years": [int(year) for year in YEAR_ORDER],
        "national": [national_counts[year] for year in YEAR_ORDER],
        "showrooms": {
            cdsid: [showroom_counts[(showroom, year)] for year in YEAR_ORDER]
            for showroom, cdsid in showroom_cdsid.items()
        },
    }
    args.sent_output.write_text(
        json.dumps(sent_payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"matched {matched_rows:,} sent rows across "
        f"{len(matched_employees):,} current employees; "
        f"latest responses for {len(latest_response):,} employees; "
        f"sent through {sent_through}"
    )


if __name__ == "__main__":
    main()
