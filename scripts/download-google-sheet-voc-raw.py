"""Download the private-field VOC source tabs into short-lived workbooks.

The generated files are intended for the GitHub Actions runner temporary
directory only.  They must never be committed: downstream scripts retain only
privacy-safe counts, score totals, dates, and reviewed keyword summaries.
"""

from __future__ import annotations

import argparse
import csv
import io
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import openpyxl


WORKBOOK_ID = "1KZust31kwsHrv0VZEqyPhACza9R3rZibOdf467c6JXA"
SENT_SHEET = "VOC발신건수"
RESPONSE_SHEET = "voc회신건수"
YEARS = ("2023", "2024", "2025", "2026")
USER_AGENT = "Mozilla/5.0 VolvoDataDashboard/1.0"


def clean(value: Any) -> str:
    return str(value or "").replace("\u00a0", " ").strip()


def parse_date(value: Any) -> datetime | None:
    text = clean(value)
    if not text:
        return None
    normalized = re.sub(r"\s+", " ", text).replace(".", "-").replace("/", "-")
    normalized = re.sub(r"-\s*", "-", normalized).strip(" -")
    for candidate in (normalized, normalized[:19], normalized[:10]):
        try:
            return datetime.fromisoformat(candidate)
        except ValueError:
            continue
    return None


def fetch_rows(sheet: str, header_rows: int) -> list[list[str]]:
    query = urlencode(
        {
            "tqx": "out:csv",
            "sheet": sheet,
            "headers": str(header_rows),
        }
    )
    url = f"https://docs.google.com/spreadsheets/d/{WORKBOOK_ID}/gviz/tq?{query}"
    request = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(request, timeout=180) as response:
        text = response.read().decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise RuntimeError(f"{sheet}: 원자료를 읽지 못했습니다.")
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def header_index(headers: list[str], required: set[str]) -> dict[str, int]:
    normalized = {clean(value).replace("\n", ""): index for index, value in enumerate(headers)}
    missing = required - set(normalized)
    if missing:
        raise RuntimeError(f"원자료 필수 열이 없습니다: {sorted(missing)}")
    return normalized


def write_sent_workbook(rows: list[list[str]], destination: Path) -> tuple[int, datetime]:
    headers = rows[0]
    columns = header_index(
        headers,
        {"전시장 방문일시", "소속 딜러명", "소속 전시장명", "담당 영업직원"},
    )
    visit_index = columns["전시장 방문일시"]
    valid_rows: list[list[Any]] = []
    latest: datetime | None = None
    for source in rows[1:]:
        visited_at = parse_date(source[visit_index])
        if visited_at is None or str(visited_at.year) not in YEARS:
            continue
        row: list[Any] = list(source)
        row[visit_index] = visited_at
        valid_rows.append(row)
        latest = visited_at if latest is None or visited_at > latest else latest
    if latest is None or len(valid_rows) < 300:
        raise RuntimeError("VOC 발신 원자료의 유효 행이 비정상적으로 적습니다.")

    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet("Raw Data")
    worksheet.append([])
    worksheet.append(headers)
    for row in valid_rows:
        worksheet.append(row)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return len(valid_rows), latest


def write_response_workbook(
    rows: list[list[str]], destination: Path
) -> tuple[Counter[str], datetime]:
    headers = rows[0]
    columns = header_index(
        headers,
        {"전시장 방문일", "방문", "상담영업직원", "상담 만족도"},
    )
    visit_index = columns["전시장 방문일"]
    score_index = columns["상담 만족도"]
    by_year: dict[str, list[list[Any]]] = {year: [] for year in YEARS}
    latest: datetime | None = None
    for source in rows[1:]:
        visited_at = parse_date(source[visit_index])
        if visited_at is None:
            continue
        year = str(visited_at.year)
        if year not in by_year:
            continue
        try:
            score = float(clean(source[score_index]).replace(",", ""))
        except ValueError:
            continue
        if not 0 <= score <= 10:
            raise RuntimeError("VOC 상담 만족도 범위를 벗어난 값이 있습니다.")
        row: list[Any] = list(source)
        row[visit_index] = visited_at
        row[score_index] = score
        by_year[year].append(row)
        latest = visited_at if latest is None or visited_at > latest else latest
    counts = Counter({year: len(values) for year, values in by_year.items()})
    if latest is None or counts["2026"] < 300:
        raise RuntimeError("VOC 회신 원자료의 2026년 유효 행이 비정상적으로 적습니다.")

    workbook = openpyxl.Workbook(write_only=True)
    for year in YEARS:
        worksheet = workbook.create_sheet(year)
        worksheet.append([])
        worksheet.append(headers)
        for row in by_year[year]:
            worksheet.append(row)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return counts, latest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sent-output", required=True, type=Path)
    parser.add_argument("--response-output", required=True, type=Path)
    args = parser.parse_args()

    sent_rows = fetch_rows(SENT_SHEET, header_rows=2)
    response_rows = fetch_rows(RESPONSE_SHEET, header_rows=2)
    sent_count, sent_latest = write_sent_workbook(sent_rows, args.sent_output)
    response_counts, response_latest = write_response_workbook(
        response_rows, args.response_output
    )
    print(
        "Google Sheet VOC 원자료 확인 완료: "
        f"발신 {sent_count:,}건({sent_latest.date().isoformat()}까지), "
        f"회신 {sum(response_counts.values()):,}건"
        f"({response_latest.date().isoformat()}까지)"
    )


if __name__ == "__main__":
    main()
