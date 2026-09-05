"""Safely reconcile the dashboard staff roster with Sales-DMS.

Only showroom, employee name, role, effective title and hire date are retained.
Employee IDs, CDSIDs, email addresses and phone numbers are never written.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "app" / "data" / "voc-staff-analysis.json"
LOGIN_URL = (
    os.environ.get("VOLVO_SALES_URL")
    or "https://sales.volvocars.kr/login/login.asp"
).strip()
TARGET_ROLES = {"영업직원", "영업팀장"}
MINIMUM_SAFE_ROSTER = 300


def text(value: Any) -> str:
    return str(value or "").strip()


def header_map(values: list[Any] | tuple[Any, ...]) -> dict[str, int]:
    return {text(value).replace("\n", ""): index for index, value in enumerate(values) if text(value)}


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    normalized = text(value).replace(".", "-").replace("/", "-")[:10]
    return date.fromisoformat(normalized)


def full_months(start: date, end: date) -> int:
    months = (end.year - start.year) * 12 + end.month - start.month
    if end.day < start.day:
        months -= 1
    return max(0, months)


def tenure_bucket(months: int) -> tuple[str, str]:
    if months <= 3:
        return "under-3m", "3개월 이내"
    if months <= 12:
        return "under-1y", "1년 이내"
    if months <= 36:
        return "1-3y", "1~3년"
    if months <= 60:
        return "3-5y", "3~5년"
    if months <= 120:
        return "5-10y", "5~10년"
    return "over-10y", "10년 이상"


def effective_job_title(role: Any, job_title: Any) -> str:
    normalized_role = text(role)
    normalized_title = text(job_title)
    return "팀장" if normalized_role == "영업팀장" else normalized_title


def normalize_row(raw: dict[str, Any]) -> dict[str, str] | None:
    role = text(raw.get("직원권한"))
    showroom = text(raw.get("전시장명") or raw.get("전시장"))
    status = text(raw.get("자동배정 여부"))
    departed = text(raw.get("퇴사일자"))
    if role not in TARGET_ROLES or showroom == "Volvo Car Korea":
        return None
    if status and status != "활성":
        return None
    if departed:
        return None
    hire_date = parse_date(raw.get("입사일자"))
    name = text(raw.get("직원명"))
    if not showroom or not name:
        raise RuntimeError("전시장명 또는 직원명이 비어 있는 Sales-DMS 행이 있습니다.")
    title = effective_job_title(role, raw.get("직급"))
    if not title:
        raise RuntimeError(f"직급이 비어 있습니다: {showroom} / {name}")
    return {
        "dmsShowroom": showroom,
        "role": role,
        "jobTitle": title,
        "name": name,
        "hireDate": hire_date.isoformat(),
    }


def load_workbook_rows(path: Path) -> tuple[list[dict[str, Any]], str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = header_map(next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True)))
    required = {"전시장명", "직원권한", "직원명", "직급", "입사일자"}
    if missing := required - set(headers):
        raise RuntimeError(f"직원 명단 필수 열이 없습니다: {sorted(missing)}")
    rows = [
        {label: row[index] for label, index in headers.items()}
        for row in worksheet.iter_rows(min_row=4, values_only=True)
    ]
    return rows, worksheet.title


def visible(locator: Any) -> bool:
    try:
        return locator.is_visible()
    except Exception:
        return False


def employee_grid(page: Any) -> tuple[Any, Any] | None:
    for candidate_page in reversed(page.context.pages):
        for frame in candidate_page.frames:
            body = frame.locator("body").inner_text(timeout=5_000)
            if frame.locator("#com_cd").count() and "직원 목록" in body and "직원 CDSID" in body:
                return candidate_page, frame
    return None


def wait_for_grid(page: Any, timeout_ms: int = 30_000) -> tuple[Any, Any]:
    deadline = datetime.now().timestamp() + timeout_ms / 1_000
    while datetime.now().timestamp() < deadline:
        if found := employee_grid(page):
            return found
        page.wait_for_timeout(500)
    raise RuntimeError("Sales-DMS 직원등록 표를 찾지 못했습니다.")


def click_menu(page: Any, label: str) -> None:
    pattern = re.compile(re.escape(label), re.IGNORECASE)
    for candidate_page in reversed(page.context.pages):
        for frame in candidate_page.frames:
            matches = frame.get_by_text(pattern)
            for index in range(matches.count()):
                match = matches.nth(index)
                if visible(match):
                    match.click(timeout=15_000)
                    candidate_page.wait_for_timeout(800)
                    return
    raise RuntimeError(f"Sales-DMS 메뉴를 찾지 못했습니다: {label}")


def login_and_open_grid(page: Any, user_id: str, password: str) -> tuple[Any, Any]:
    page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=60_000)
    boxes = page.get_by_role("textbox")
    if boxes.count() != 2:
        raise RuntimeError("Sales-DMS 로그인 입력창 구조가 변경되었습니다.")
    boxes.nth(0).fill(user_id)
    boxes.nth(1).fill(password)
    page.get_by_role("button", name="Login", exact=True).click()
    page.wait_for_load_state("domcontentloaded", timeout=60_000)
    page.wait_for_timeout(3_000)
    if "/login/" in page.url.lower():
        raise RuntimeError("Sales-DMS 로그인에 실패했습니다. GitHub Secret을 확인해 주세요.")
    if found := employee_grid(page):
        return found
    for label in ("Master data management", "Master 관리", "사용자 등록", "직원등록"):
        click_menu(page, label)
    return wait_for_grid(page)


def table_rows(frame: Any, expected_role: str) -> list[dict[str, str]]:
    rows = frame.locator("tr").evaluate_all(
        """rows => rows.map(row => Array.from(row.querySelectorAll('th,td'))
        .map(cell => (cell.textContent || '').trim().replace(/\\s+/g, ' ')))"""
    )
    header_index = next((i for i, row in enumerate(rows) if "직원 CDSID" in row and "직원권한" in row), -1)
    if header_index < 0:
        raise RuntimeError("Sales-DMS 직원 목록 열 구조가 변경되었습니다.")
    headers = rows[header_index]
    result: list[dict[str, str]] = []
    for values in rows[header_index + 1 :]:
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        if text(record.get("직원권한")) == expected_role:
            result.append(record)
    return result


def set_page_size(page: Any) -> None:
    for frame in page.frames:
        label = frame.get_by_text(re.compile(r"리스트\s*갯수", re.IGNORECASE))
        if not label.count():
            continue
        input_box = label.first.locator("xpath=parent::*").locator("input:not([type='hidden'])")
        if input_box.count() and visible(input_box.first):
            input_box.first.fill("500")
            input_box.first.press("Enter")
            wait_for_grid(page)
            return
    raise RuntimeError("Sales-DMS 직원 목록 표시 건수 입력창을 찾지 못했습니다.")


def collect_dms_rows(user_id: str, password: str) -> list[dict[str, str]]:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page()
        try:
            _, frame = login_and_open_grid(page, user_id, password)
            status = frame.locator("select[name='resign_type']")
            if status.count():
                status.select_option(label="재직자")
            for selector in ("#s_date", "#e_date"):
                if frame.locator(selector).count():
                    frame.locator(selector).fill("")
            set_page_size(page)
            collected: list[dict[str, str]] = []
            for role in sorted(TARGET_ROLES):
                _, frame = wait_for_grid(page)
                frame.locator("#com_cd").select_option(label=role)
                search = frame.locator("#search")
                search.click(timeout=10_000)
                page.wait_for_timeout(1_000)
                _, frame = wait_for_grid(page)
                collected.extend(table_rows(frame, role))
            return collected
        finally:
            browser.close()


def default_history() -> dict[str, Any]:
    return {
        "years": {},
        "commentResponses": 0,
        "strengthKeywords": [],
        "improvementKeywords": [],
    }


def reconcile(payload: dict[str, Any], raw_rows: list[dict[str, Any]], as_of: date) -> dict[str, int]:
    showroom_key_by_dms = {
        text(showroom["dmsShowroom"]): key for key, showroom in payload["showrooms"].items()
    }
    normalized = [item for row in raw_rows if (item := normalize_row(row)) is not None]
    normalized = [item for item in normalized if item["dmsShowroom"] in showroom_key_by_dms]
    keys = [(item["dmsShowroom"], item["name"]) for item in normalized]
    if len(keys) != len(set(keys)):
        raise RuntimeError("Sales-DMS 명단에 전시장·직원명 중복이 있습니다.")
    if len(normalized) < MINIMUM_SAFE_ROSTER:
        raise RuntimeError(f"명단이 비정상적으로 적어 동기화를 중단했습니다: {len(normalized)}명")
    found_showrooms = {item["dmsShowroom"] for item in normalized}
    expected_showrooms = set(showroom_key_by_dms)
    if found_showrooms != expected_showrooms:
        raise RuntimeError(
            f"39개 전시장 불일치: 누락={sorted(expected_showrooms - found_showrooms)}, "
            f"초과={sorted(found_showrooms - expected_showrooms)}"
        )

    incoming: dict[str, list[dict[str, str]]] = {key: [] for key in payload["showrooms"]}
    for item in normalized:
        incoming[showroom_key_by_dms[item["dmsShowroom"]]].append(item)

    added = removed = updated = 0
    for showroom_key, showroom in payload["showrooms"].items():
        current = {text(item["name"]): item for item in showroom.get("employees", [])}
        former = {text(item["name"]): item for item in showroom.get("formerEmployees", [])}
        next_employees: list[dict[str, Any]] = []
        active_names: set[str] = set()
        for item in incoming[showroom_key]:
            active_names.add(item["name"])
            previous = current.get(item["name"]) or former.get(item["name"])
            if previous is None:
                added += 1
                employee: dict[str, Any] = {**default_history()}
            else:
                employee = {
                    key: value
                    for key, value in previous.items()
                    if key not in {"archivedAt", "archiveReason"}
                }
            before = {key: employee.get(key) for key in ("role", "jobTitle", "hireDate")}
            months = full_months(date.fromisoformat(item["hireDate"]), as_of)
            bucket_id, bucket_label = tenure_bucket(months)
            employee.update(
                {
                    "name": item["name"],
                    "role": item["role"],
                    "jobTitle": item["jobTitle"],
                    "hireDate": item["hireDate"],
                    "tenureMonths": months,
                    "tenureBucket": bucket_id,
                    "tenureBucketLabel": bucket_label,
                }
            )
            if previous is not None and before != {key: employee.get(key) for key in before}:
                updated += 1
            next_employees.append(employee)

        archived = [item for name, item in former.items() if name not in active_names]
        for name, employee in current.items():
            if name in active_names:
                continue
            removed += 1
            archived.append(
                {**employee, "archivedAt": as_of.isoformat(), "archiveReason": "Sales-DMS 현재 명단 제외"}
            )
        showroom["employees"] = sorted(next_employees, key=lambda item: (item["role"] != "영업직원", item["name"]))
        showroom["formerEmployees"] = sorted(archived, key=lambda item: item["name"])

    payload["source"].update(
        {
            "rosterCheckedAt": as_of.isoformat(),
            "rosterRule": "현재 재직자 · 39개 전시장 · 영업직원/영업팀장 · 직원권한 우선",
            "rosterUpdateSchedule": "매일 06:00 KST · 1일 1회",
            "jobTitleSourceDate": as_of.isoformat(),
        }
    )
    return {"employees": len(normalized), "added": added, "removed": removed, "updated": updated}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--roster", type=Path)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--as-of", default=datetime.now(ZoneInfo("Asia/Seoul")).date().isoformat())
    args = parser.parse_args()
    as_of = date.fromisoformat(args.as_of)
    payload = json.loads(args.output.read_text(encoding="utf-8"))
    if args.roster:
        rows, sheet = load_workbook_rows(args.roster)
        source_name = args.roster.name
    else:
        user_id = os.environ.get("VOLVO_SALES_ID", "").strip()
        password = os.environ.get("VOLVO_SALES_PASSWORD", "").strip()
        if not user_id or not password:
            raise RuntimeError("VOLVO_SALES_ID와 VOLVO_SALES_PASSWORD Secret이 필요합니다.")
        rows = collect_dms_rows(user_id, password)
        sheet = "Sales-DMS 직원등록"
        source_name = "Sales-DMS 직원등록"
    summary = reconcile(payload, rows, as_of)
    payload["source"].update(
        {"jobTitleSourceWorkbook": source_name, "jobTitleSourceSheet": sheet}
    )
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        "Sales-DMS 명단 동기화 완료: "
        f"{summary['employees']}명, 신규 {summary['added']}명, "
        f"변경 {summary['updated']}명, 이전 명단 {summary['removed']}명"
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"Sales-DMS 명단 동기화 실패: {error}", file=sys.stderr)
        raise SystemExit(1)
