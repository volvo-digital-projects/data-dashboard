from __future__ import annotations

import argparse
import json
import runpy
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
GENERATOR = runpy.run_path(str(Path(__file__).with_name("generate-voc-staff-analysis.py")))
YEARS = tuple(GENERATOR["YEARS"])
STRENGTH_PATTERNS = GENERATOR["STRENGTH_PATTERNS"]
IMPROVEMENT_PATTERNS = GENERATOR["IMPROVEMENT_PATTERNS"]
header_map = GENERATOR["header_map"]
keyword_summary = GENERATOR["keyword_summary"]
load_voc = GENERATOR["load_voc"]
normalise_showroom_name = GENERATOR["normalise_showroom_name"]
merged_staff_comments = GENERATOR["merged_staff_comments"]
merged_staff_metrics = GENERATOR["merged_staff_metrics"]
staff_source_keys = GENERATOR["staff_source_keys"]
current_staff_source_keys = GENERATOR["current_staff_source_keys"]


def date_value(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def load_2026_controls(
    path: Path,
    rate_cutoff: date | None,
) -> tuple[str, dict[tuple[str, str], str], dict[tuple[str, str], int]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["2026"]
    headers = header_map(next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    showroom_index = headers.get("방문 전시장", headers.get("방문"))
    name_index = headers["상담영업직원"]
    score_index = headers["상담 만족도"]
    visit_index = headers["전시장 방문일"]
    latest_date: datetime | None = None
    latest_by_employee: dict[tuple[str, str], datetime] = {}
    rate_responses: dict[tuple[str, str], int] = defaultdict(int)
    blank_rows = 0
    for row in worksheet.iter_rows(min_row=3, values_only=True):
        showroom = row[showroom_index] if showroom_index is not None else None
        name = row[name_index]
        score = row[score_index]
        if showroom is None and name is None and score is None:
            blank_rows += 1
            if blank_rows >= 250:
                break
            continue
        blank_rows = 0
        if not isinstance(score, (int, float)):
            continue
        visited_at = date_value(row[visit_index])
        if visited_at is None:
            continue
        latest_date = visited_at if latest_date is None or visited_at > latest_date else latest_date
        if showroom and name:
            key = (normalise_showroom_name(showroom), str(name).strip())
            current = latest_by_employee.get(key)
            if current is None or visited_at > current:
                latest_by_employee[key] = visited_at
            if rate_cutoff is not None and visited_at.date() <= rate_cutoff:
                rate_responses[key] += 1
    if latest_date is None:
        raise ValueError("2026 VOC 유효 회신의 전시장 방문일을 찾지 못했습니다.")
    return (
        latest_date.date().isoformat(),
        {key: value.date().isoformat() for key, value in latest_by_employee.items()},
        dict(rate_responses),
    )


def refresh_tenure_cohorts(payload: dict[str, Any]) -> None:
    cohorts = {
        item["id"]: {
            **item,
            "employeeCount": 0,
            "respondingEmployees": 0,
            "responses": 0,
            "scoreSum": 0.0,
            "average": None,
        }
        for item in payload["tenureCohorts"]
    }
    for showroom in payload["showrooms"].values():
        for employee in showroom["employees"]:
            cohort = cohorts[employee["tenureBucket"]]
            cohort["employeeCount"] += 1
            responses = sum(int(item.get("responses", 0)) for item in employee["years"].values())
            score_sum = sum(float(item.get("scoreSum", 0)) for item in employee["years"].values())
            if responses:
                cohort["respondingEmployees"] += 1
                cohort["responses"] += responses
                cohort["scoreSum"] += score_sum
    refreshed = []
    for original in payload["tenureCohorts"]:
        cohort = cohorts[original["id"]]
        cohort["scoreSum"] = round(cohort["scoreSum"], 1)
        cohort["average"] = round(cohort["scoreSum"] / cohort["responses"], 2) if cohort["responses"] else None
        refreshed.append(cohort)
    payload["tenureCohorts"] = refreshed


def refresh_staff_analysis(path: Path, output: Path) -> dict[str, Any]:
    payload = json.loads(output.read_text(encoding="utf-8"))
    national, staff_metrics, staff_comments = load_voc(path)
    sent_through = payload["source"].get("sentThrough")
    rate_cutoff = date.fromisoformat(sent_through) if sent_through else None
    through, latest_by_employee, rate_responses = load_2026_controls(path, rate_cutoff)

    responding_2026 = {
        key for key, values in staff_metrics.items() if values["2026"]["responses"]
    }
    national_responses = int(national["2026"]["responses"])
    national_score_sum = round(float(national["2026"]["scoreSum"]), 1)
    payload["nationalYears"]["2026"] = {
        "responses": national_responses,
        "scoreSum": national_score_sum,
        "average": round(national_score_sum / national_responses, 3) if national_responses else None,
        "respondingEmployees": len(responding_2026),
        "averageResponsesPerEmployee": round(national_responses / len(responding_2026), 1) if responding_2026 else None,
    }

    refreshed_employees = 0
    matched_responses = 0
    current_name_counts = Counter(
        str(employee["name"]).strip()
        for showroom in payload["showrooms"].values()
        for employee in showroom["employees"]
    )
    available_metric_keys = tuple(staff_metrics)
    claimed_source_keys = {
        source_key
        for showroom in payload["showrooms"].values()
        for employee in showroom["employees"]
        for source_key in current_staff_source_keys(
            showroom["showroom"],
            employee["name"],
            available_metric_keys,
            current_name_counts,
        )
    }
    for showroom in payload["showrooms"].values():
        showroom_name = normalise_showroom_name(showroom["showroom"])
        for employee in showroom["employees"]:
            key = (showroom_name, str(employee["name"]).strip())
            source_keys = current_staff_source_keys(
                *key,
                available_metric_keys,
                current_name_counts,
            )
            refreshed_years = merged_staff_metrics(staff_metrics, source_keys)
            for year in YEARS:
                refreshed_year = dict(refreshed_years[year])
                previous_year = employee["years"].get(year, {})
                if not refreshed_year.get("responses"):
                    if "sent" in previous_year:
                        employee["years"][year] = {
                            "responses": 0,
                            "scoreSum": 0,
                            "sent": previous_year["sent"],
                        }
                    else:
                        employee["years"].pop(year, None)
                    continue
                if "sent" in previous_year:
                    refreshed_year["sent"] = previous_year["sent"]
                if year == "2026" and "sent" in previous_year:
                    refreshed_year["rateResponses"] = sum(
                        rate_responses.get(source_key, 0)
                        for source_key in source_keys
                    )
                elif "rateResponses" in previous_year:
                    refreshed_year["rateResponses"] = previous_year["rateResponses"]
                employee["years"][year] = refreshed_year
                if year == "2026":
                    matched_responses += int(refreshed_year["responses"])

            latest = max(
                (
                    latest_by_employee[source_key]
                    for source_key in source_keys
                    if source_key in latest_by_employee
                ),
                default=None,
            )
            if latest:
                employee["latestResponseDate"] = latest
            elif employee.get("latestResponseDate", "").startswith("2026-"):
                employee.pop("latestResponseDate", None)

            comments = merged_staff_comments(staff_comments, source_keys)
            employee["commentResponses"] = len(comments)
            employee["strengthKeywords"] = keyword_summary(
                comments,
                STRENGTH_PATTERNS,
                len(STRENGTH_PATTERNS),
                exclude_negative_context=True,
            )
            employee["improvementKeywords"] = keyword_summary(
                comments,
                IMPROVEMENT_PATTERNS,
                len(IMPROVEMENT_PATTERNS),
            )
            refreshed_employees += 1

        target_names = {str(employee["name"]).strip() for employee in showroom["employees"]}
        excluded_counts: dict[str, int] = defaultdict(int)
        for (raw_showroom, name), by_year in staff_metrics.items():
            source_key = (normalise_showroom_name(raw_showroom), name)
            if (
                source_key[0] != showroom_name
                or name in target_names
                or source_key in claimed_source_keys
            ):
                continue
            excluded_counts[name] += sum(int(values["responses"]) for values in by_year.values())
        showroom["excludedRawNames"] = [
            {
                "name": name,
                "responses": responses,
                "reason": "현재 DMS 재직 명단 미확인",
            }
            for name, responses in sorted(
                excluded_counts.items(), key=lambda item: (-item[1], item[0])
            )
        ]

    refresh_tenure_cohorts(payload)
    source = payload["source"]
    source["workbook"] = path.name
    source["vocThrough"] = through
    source["latestResponseThrough"] = through
    source["commentAnalysis"] = (
        "2023-2026 YTD 원문 문장별 긍정·부정 맥락 분리 · "
        f"{len(STRENGTH_PATTERNS)}개 강점/{len(IMPROVEMENT_PATTERNS)}개 보완 주제"
    )
    source["commentAnalysisEmployees"] = refreshed_employees
    source["refreshScope"] = (
        "기존 2023-2025 집계 보존 · 2026 점수/회신/원문 분석 갱신 · "
        "이름이 유일한 재직 직원의 전시장 이동 이력 자동 통합"
    )
    source["staffHistoryMatchRule"] = (
        "현재 재직 명단에서 이름이 유일한 직원은 2023-2026 전시장 이동 이력 통합 · "
        "동명이인은 전시장 기준 분리"
    )
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"staff analysis: 2026 national {national_responses:,} responses; "
        f"current-roster matched {matched_responses:,}; through {through}"
    )
    return payload


def average_and_count(score_sum: float, count: int) -> tuple[float | None, int]:
    return (round(score_sum / count, 6), count) if count else (None, 0)


def refresh_consultation(path: Path, output: Path, staff_payload: dict[str, Any]) -> None:
    consultation = json.loads(output.read_text(encoding="utf-8"))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["2026"]
    headers = header_map(next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    showroom_index = headers.get("방문 전시장", headers.get("방문"))
    score_index = headers["상담 만족도"]
    visit_index = headers["전시장 방문일"]
    rate_cutoff_text = staff_payload["source"].get("sentThrough")
    rate_cutoff = date.fromisoformat(rate_cutoff_text) if rate_cutoff_text else None
    national_score_sum = 0.0
    national_count = 0
    national_rate_count = 0
    showroom_totals: dict[str, list[float | int]] = defaultdict(lambda: [0.0, 0])
    showroom_rate_counts: dict[str, int] = defaultdict(int)
    blank_rows = 0
    for row in worksheet.iter_rows(min_row=3, values_only=True):
        showroom = row[showroom_index] if showroom_index is not None else None
        score = row[score_index]
        if showroom is None and score is None:
            blank_rows += 1
            if blank_rows >= 250:
                break
            continue
        blank_rows = 0
        if not isinstance(score, (int, float)):
            continue
        national_score_sum += float(score)
        national_count += 1
        visited_at = date_value(row[visit_index])
        within_rate_cutoff = (
            rate_cutoff is not None
            and visited_at is not None
            and visited_at.date() <= rate_cutoff
        )
        if within_rate_cutoff:
            national_rate_count += 1
        if showroom:
            key = normalise_showroom_name(showroom)
            showroom_totals[key][0] = float(showroom_totals[key][0]) + float(score)
            showroom_totals[key][1] = int(showroom_totals[key][1]) + 1
            if within_rate_cutoff:
                showroom_rate_counts[key] += 1

    showroom_by_name = {
        normalise_showroom_name(item["showroom"]): cdsid
        for cdsid, item in staff_payload["showrooms"].items()
    }
    if set(showroom_totals) != set(showroom_by_name):
        raise ValueError(
            "2026 VOC 전시장 매핑 불일치: "
            f"누락={sorted(set(showroom_by_name) - set(showroom_totals))}, "
            f"초과={sorted(set(showroom_totals) - set(showroom_by_name))}"
        )

    def replace_2026(values: list[Any], average: float | None, responses: int) -> None:
        year_index = consultation["years"].index(2026)
        values[2 + year_index * 2] = average
        values[3 + year_index * 2] = responses
        weighted_sum = sum(
            float(values[2 + index * 2]) * int(values[3 + index * 2])
            for index in range(len(consultation["years"]))
            if values[2 + index * 2] is not None and int(values[3 + index * 2])
        )
        total_responses = sum(int(values[3 + index * 2]) for index in range(len(consultation["years"])))
        values[0] = round(weighted_sum / total_responses, 6) if total_responses else None
        values[1] = total_responses

    national_average, national_responses = average_and_count(national_score_sum, national_count)
    replace_2026(consultation["national"], national_average, national_responses)
    rate_years = consultation.setdefault(
        "responseRateResponses",
        {
            "national": [consultation["national"][3 + index * 2] for index in range(len(consultation["years"]))],
            "showrooms": {
                cdsid: [values[3 + index * 2] for index in range(len(consultation["years"]))]
                for cdsid, values in consultation["showrooms"].items()
            },
        },
    )
    year_index = consultation["years"].index(2026)
    rate_years["national"][year_index] = national_rate_count
    for showroom_name, cdsid in showroom_by_name.items():
        score_sum, responses = showroom_totals[showroom_name]
        average, count = average_and_count(float(score_sum), int(responses))
        replace_2026(consultation["showrooms"][cdsid], average, count)
        rate_years["showrooms"][cdsid][year_index] = showroom_rate_counts[showroom_name]

    consultation["updatedThrough"] = staff_payload["source"]["vocThrough"]
    consultation["sourceWorkbook"] = path.name
    consultation["refreshScope"] = "2026 원본 회신만 갱신 · 2023-2025 기존 집계 보존"
    consultation["responseRateThrough"] = rate_cutoff_text
    output.write_text(json.dumps(consultation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"consultation history: 2026 national {national_count:,} responses")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--voc", required=True, type=Path)
    parser.add_argument(
        "--staff-output",
        type=Path,
        default=ROOT / "app" / "data" / "voc-staff-analysis.json",
    )
    parser.add_argument(
        "--consultation-output",
        type=Path,
        default=ROOT / "app" / "data" / "voc-consultation.json",
    )
    args = parser.parse_args()
    staff_payload = refresh_staff_analysis(args.voc, args.staff_output)
    refresh_consultation(args.voc, args.consultation_output, staff_payload)


if __name__ == "__main__":
    main()
